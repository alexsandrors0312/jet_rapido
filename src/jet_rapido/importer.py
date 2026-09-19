"""Importa o layout SPX preservando a origem e sem executar células."""

from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from hashlib import sha256
from math import isfinite
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


HEADERS = (
    "AT ID", "Sequence", "Stop", "SPX TN", "Destination Address",
    "Bairro", "City", "Zipcode/Postal code", "Latitude", "Longitude",
)


class ImportValidationError(ValueError):
    """Erro recuperável de contrato de importação."""


@dataclass(frozen=True)
class Package:
    source_row: int
    route_id: str
    tracking_id: str
    original_sequence: int | None
    original_stop: int | None
    address: str
    neighborhood: str
    city: str
    postal_code: str
    latitude: float
    longitude: float
    street_key: str


def normalized_text(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(c for c in value if not unicodedata.combining(c))
    return " ".join(value.casefold().split())


def street_key(address: str, city: str) -> str:
    # Chave textual candidata. Não é identidade de trecho nem geocodificação.
    street = normalized_text(address.split(",", 1)[0])
    street = re.sub(r"^(?:r\.?|rua)\s+", "rua ", street)
    street = re.sub(r"^(?:av\.?|avenida)\s+", "avenida ", street)
    street = re.sub(r"^(?:tv\.?|travessa)\s+", "travessa ", street)
    return f"{normalized_text(city)}|{street}"


def _text(value, field: str, row: int, max_length: int) -> str:
    if not isinstance(value, str) or not value.strip() or value.strip() == "-":
        raise ImportValidationError(f"Linha {row}: {field} deve ser texto preenchido.")
    value = value.strip()
    if len(value) > max_length:
        raise ImportValidationError(f"Linha {row}: {field} excede {max_length} caracteres.")
    return value


def _position(value, field: str, row: int) -> int | None:
    if value is None or (isinstance(value, str) and value.strip() in ("", "-")):
        return None
    try:
        number = float(value)
        if isinstance(value, bool) or not isfinite(number) or number <= 0 or not number.is_integer():
            raise ValueError
        return int(number)
    except (TypeError, ValueError):
        raise ImportValidationError(f"Linha {row}: {field} deve ser inteiro positivo ou vazio.") from None


def _coordinate(value, limit: int, field: str, row: int) -> float:
    try:
        if isinstance(value, bool):
            raise ValueError
        number = float(value.replace(",", ".") if isinstance(value, str) else value)
        if not isfinite(number) or not -limit <= number <= limit:
            raise ValueError
        return number
    except (TypeError, ValueError):
        raise ImportValidationError(f"Linha {row}: {field} inválida.") from None


def import_workbook(path: str | Path, sheet: str | None = None) -> dict:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ImportValidationError("Nesta etapa, somente arquivos .xlsx são aceitos.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet is None and len(workbook.sheetnames) != 1:
            raise ImportValidationError("Há várias abas. Informe --sheet explicitamente.")
        if sheet is not None and sheet not in workbook.sheetnames:
            raise ImportValidationError("Aba não encontrada.")
        ws = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        declared_dimension = ws.calculate_dimension()
        # A exportação real declara A1:A1, embora possua A1:J33.
        ws.reset_dimensions()
        rows = ws.iter_rows()
        first = next(rows, ())
        names = [str(cell.value).strip() if cell.value is not None else "" for cell in first]
        if any(cell.data_type == "f" for cell in first):
            raise ImportValidationError("Cabeçalho não pode conter fórmulas.")
        if len(names) != len(set(names)) or any(h not in names for h in HEADERS):
            raise ImportValidationError("Cabeçalhos ausentes ou duplicados. Consulte docs/IMPORTACAO.md.")
        packages = []
        warnings = []
        seen = set()
        for row_number, cells in enumerate(rows, start=2):
            if all(c.value is None for c in cells):
                continue
            if any(c.data_type in ("f", "e") for c in cells):
                raise ImportValidationError(f"Linha {row_number}: fórmula ou erro de Excel não permitido.")
            values = dict(zip(names, [c.value for c in cells]))
            limits = {
                "AT ID": 255,
                "SPX TN": 255,
                "Destination Address": 1000,
                "Bairro": 255,
                "City": 255,
                "Zipcode/Postal code": 32,
            }
            t = {h: _text(values.get(h), h, row_number, limit) for h, limit in limits.items()}
            if t["SPX TN"] in seen:
                raise ImportValidationError(f"Linha {row_number}: identificador de pacote duplicado.")
            seen.add(t["SPX TN"])
            candidate_street_key = street_key(t["Destination Address"], t["City"])
            if len(candidate_street_key) > 512:
                raise ImportValidationError(f"Linha {row_number}: chave de rua excede 512 caracteres.")
            package = Package(
                source_row=row_number, route_id=t["AT ID"], tracking_id=t["SPX TN"],
                original_sequence=_position(values.get("Sequence"), "Sequence", row_number),
                original_stop=_position(values.get("Stop"), "Stop", row_number),
                address=t["Destination Address"], neighborhood=t["Bairro"], city=t["City"],
                postal_code=t["Zipcode/Postal code"],
                latitude=_coordinate(values.get("Latitude"), 90, "Latitude", row_number),
                longitude=_coordinate(values.get("Longitude"), 180, "Longitude", row_number),
                street_key=candidate_street_key,
            )
            if package.original_stop is None or package.original_sequence is None:
                warnings.append({"row": row_number, "code": "MISSING_ORIGINAL_ORDER"})
            if re.search(r"\bs\s*/\s*n\b", package.address, re.IGNORECASE):
                warnings.append({"row": row_number, "code": "ADDRESS_WITHOUT_NUMBER"})
            if not re.fullmatch(r"\d{5}-?\d{3}", package.postal_code):
                warnings.append({"row": row_number, "code": "POSTAL_CODE_REVIEW"})
            if (package.latitude, package.longitude) == (0, 0):
                warnings.append({"row": row_number, "code": "COORDINATE_REVIEW"})
            packages.append(package)
        if not packages:
            raise ImportValidationError("Nenhum pacote encontrado.")
        groups = defaultdict(list)
        for p in packages:
            groups[(p.route_id, p.street_key)].append(p)
        split_streets = []
        for (_, key), members in sorted(groups.items()):
            stops = sorted({p.original_stop for p in members if p.original_stop is not None})
            if len(stops) > 1:
                split_streets.append({
                    "route_id": members[0].route_id,
                    "street_key": key, "original_stops": stops,
                    "source_rows": [p.source_row for p in members],
                    "package_count": len(members),
                })
        coordinates = Counter((p.latitude, p.longitude) for p in packages)
        return {
            "schema_version": 1,
            "source": {"sha256": sha256(path.read_bytes()).hexdigest(),
                       "sheet": ws.title, "declared_dimension": declared_dimension},
            "summary": {
                "packages": len(packages),
                "routes": len({p.route_id for p in packages}),
                "original_stops": len({(p.route_id, p.original_stop) for p in packages if p.original_stop is not None}),
                "without_stop": sum(p.original_stop is None for p in packages),
                "without_sequence": sum(p.original_sequence is None for p in packages),
                "unique_coordinates": len(coordinates),
                "shared_coordinate_groups": sum(v > 1 for v in coordinates.values()),
                "split_street_candidates": len(split_streets),
            },
            "warnings": warnings,
            "split_streets": split_streets,
            "packages": [asdict(p) for p in packages],
        }
    finally:
        workbook.close()
